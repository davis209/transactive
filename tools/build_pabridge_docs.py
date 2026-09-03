from __future__ import print_function

import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
TABLE_FILL = "E8EEF5"
CODE_FILL = "F3F5F7"
MUTED = "6B7280"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    tbl_pr.append(indent)
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.add_run("PABridgeAgent  |  ").font.color.rgb = RGBColor.from_string(MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def configure_document(doc, title):
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, color, before, after in (
        ("Title", 22, BLUE, 0, 10),
        ("Heading 1", 15, BLUE, 15, 7),
        ("Heading 2", 12, BLUE, 11, 5),
        ("Heading 3", 10.5, DARK_BLUE, 8, 3),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    code.font.size = Pt(8)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(4)
    code.paragraph_format.left_indent = Inches(0.15)
    code.paragraph_format.right_indent = Inches(0.15)
    code.paragraph_format.line_spacing = 1.0

    header = section.header.paragraphs[0]
    header.text = title
    header.style = styles["Normal"]
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    footer = section.footer.paragraphs[0]
    add_page_number(footer)


def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CODE_FILL)
    p = cell.paragraphs[0]
    p.style = "Code Block"
    p.add_run("\n".join(lines))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def column_widths(count):
    if count == 2:
        return [3000, 6360]
    if count == 3:
        return [2200, 1800, 5360]
    if count == 4:
        return [1800, 1300, 1100, 5160]
    base = 9360 // count
    return [base] * (count - 1) + [9360 - base * (count - 1)]


def add_markdown_table(doc, rows):
    cells = [[item.strip() for item in row.strip().strip("|").split("|")] for row in rows]
    cells = [row for row in cells if not all(re.match(r"^:?-{3,}:?$", item) for item in row)]
    if not cells:
        return
    count = max(len(row) for row in cells)
    table = doc.add_table(rows=len(cells), cols=count)
    table.style = "Table Grid"
    set_table_geometry(table, column_widths(count))
    for r, values in enumerate(cells):
        for c in range(count):
            cell = table.cell(r, c)
            text = values[c] if c < len(values) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text.replace("`", ""))
            run.font.size = Pt(8.5)
            if r == 0:
                set_cell_shading(cell, TABLE_FILL)
                run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_text_paragraph(doc, line):
    line = re.sub(r"^>\s?", "", line)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", line)
    for part in parts:
        if not part:
            continue
        run = p.add_run(part.strip("`").strip("*"))
        if part.startswith("`"):
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
        elif part.startswith("**"):
            run.bold = True


def build_docx(source, target):
    with open(source, "r", encoding="utf-8") as handle:
        lines = handle.read().splitlines()
    title = next((line[2:].strip() for line in lines if line.startswith("# ")), os.path.basename(source))
    doc = Document()
    configure_document(doc, title)
    i = 0
    first_title = True
    while i < len(lines):
        line = lines[i]
        if not line.strip() or line.strip() == "---":
            i += 1
            continue
        if line.startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            add_code_block(doc, code)
        elif re.match(r"^#{1,3}\s+", line):
            level = len(line) - len(line.lstrip("#"))
            text = line[level:].strip()
            if level == 1 and first_title:
                p = doc.add_paragraph(style="Title")
                p.add_run(text)
                first_title = False
            else:
                doc.add_paragraph(text, style="Heading %d" % level)
        elif line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_markdown_table(doc, table_lines)
            continue
        elif re.match(r"^[-*]\s+", line):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            p.add_run(re.sub(r"^[-*]\s+", "", line).replace("`", ""))
        elif re.match(r"^\d+\.\s+", line):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(2)
            p.add_run(re.sub(r"^\d+\.\s+", "", line).replace("`", ""))
        else:
            add_text_paragraph(doc, line)
        i += 1
    doc.save(target)


def main():
    if len(sys.argv) != 2:
        raise SystemExit("Usage: build_pabridge_docs.py <docs_directory>")
    directory = sys.argv[1]
    pairs = (
        ("PABridgeAgent_Kafka_Integration.md", "PABridgeAgent_Kafka_Integration.docx"),
        ("PABridgeAgent_Kafka_Integration_EN.md", "PABridgeAgent_Kafka_Integration_EN.docx"),
        ("PABridgeAgent_REST_API.md", "PABridgeAgent_REST_API.docx"),
        ("PABridgeAgent_REST_API_EN.md", "PABridgeAgent_REST_API_EN.docx"),
    )
    for source, target in pairs:
        build_docx(os.path.join(directory, source), os.path.join(directory, target))
        print(target)


if __name__ == "__main__":
    main()
